#!/usr/bin/env python
"""Rouo PPT · 水彩风交付装配器（HTML 设计稿 → 原生可编辑 PPTX）

把 references/27-pptx-assembly.md 的整套方法做成可复用工具，四个子命令：

    optimize <bake_dir>
        烘图后处理。约定 bake_dir 下有 paper/ 与 wash/ 两个子目录：
          paper/*.png -> *.jpg        （不透明纸纹底，JPEG 存得住）
          wash/*.png  -> 原地量化      （透明晕染层，256 色 + alpha）

    assemble <spec.json> <out.pptx>
        按规格装配。规格格式见下方 SPEC FORMAT。

    measure <ref_dir> <pptx_png_dir> <regions.json>
        逐元素量墨水包围盒，输出 HTML 基准与 PPTX 回导之间的 dx/dy。
        用于回填 spec 里的 "dy"。

    export <pptx> <out_dir> [width]
        用 PowerPoint COM 把 PPTX 逐页导出 PNG（验收用）。需要 pywin32。

--------------------------------------------------------------------------
SPEC FORMAT（JSON，UTF-8）
--------------------------------------------------------------------------
{
  "width_px": 1920, "height_px": 1080,      // 设计稿尺寸（16:9）
  "paper_color": "F7F0E4",                  // 幻灯片底色，兜底用
  "dy": { "s1.assert": 7, "s3.figure": 12 },// 逐元素垂直微调（px），由 measure 回填
  "slides": [
    {
      "paper": "bake/paper/C-1-cover.jpg",  // 相对 spec 文件所在目录
      "wash":  "bake/wash/C-1-cover.png",
      "blocks": [
        {
          "key": "s1.assert",               // 与 dy 的键对应
          "left": 226, "top": 672, "width": 1400, "height": 320,
          "size_px": 96, "line_height": 1.42,
          "line_px": null,                  // 可选：显式行距(px)。小字号列必须显式继承行的统一行距
          "latin": "Noto Serif SC", "ea": "Noto Serif SC",
          "spc": 48,                        // 字距，单位百分之一磅（.01em@96px -> 48）
          "italic": false,
          "paras": [                        // 每个数组 = 一段（对应 HTML 的 <br>）
            [["十月一日，", "241F1A"]],
            [["一个国家的", "241F1A"], ["生日", "A8332C"], ["。", "241F1A"]]
          ]                                 // run 可选第三项：{"size_px":..,"spc":..,"latin":..}
        }
      ]
    }
  ]
}
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

# 幻灯片固定 13.333in × 7.5in（16:9），与设计稿宽度做精确整除映射。
# 1920px 时：1px = 6350 EMU = 0.5pt —— 都是整数，不要改用 96DPI 那套。
SLIDE_W_EMU = 12192000
SLIDE_H_EMU = 6858000
SLIDE_W_PT = 960.0


# ===========================================================================
# optimize · 烘图后处理
# ===========================================================================

def cmd_optimize(argv: list[str]) -> int:
    if len(argv) != 1:
        print("用法：optimize <bake_dir>", file=sys.stderr)
        return 2
    from PIL import Image

    bake = Path(argv[0])
    total = 0.0
    for png in sorted((bake / "paper").glob("*.png")):
        jpg = png.with_suffix(".jpg")
        Image.open(png).convert("RGB").save(
            jpg, "JPEG", quality=94, optimize=True, subsampling=0)
        png.unlink()
        mb = jpg.stat().st_size / 1048576
        total += mb
        print(f"paper/{jpg.name:24s} {mb:5.2f} MB (JPEG q94)")

    for png in sorted((bake / "wash").glob("*.png")):
        im = Image.open(png).convert("RGBA")
        # 256 色足够：调色板里同时容纳颜色与 alpha 层级。
        # 实测 PSNR 47–54 dB（>40dB 在可察觉阈值以下），4 倍放大毛边无色带。
        im.quantize(colors=256, method=Image.FASTOCTREE,
                    dither=Image.FLOYDSTEINBERG).save(png, "PNG", optimize=True)
        mb = png.stat().st_size / 1048576
        total += mb
        print(f"wash/{png.name:25s} {mb:5.2f} MB (PNG-8/alpha)")

    print(f"\n烘图合计 {total:.2f} MB")
    return 0


# ===========================================================================
# assemble · 装配
# ===========================================================================

def _emu(px: float, per_px: float) -> int:
    return int(round(px * per_px))


def _rgb(value):
    """颜色：接受 'RRGGBB' 字符串、[r,g,b] 三元组，也接受 python-pptx 的 RGBColor。

    注意 RGBColor 是 tuple 子类，不是 int 子类 —— 想拿十六进制要用 str()，
    不要用 int()，也不要直接 json 序列化（会变成 [36,31,26] 这样的数组）。
    """
    from pptx.dml.color import RGBColor

    if isinstance(value, str):
        return RGBColor.from_string(value.lstrip("#").upper())
    if isinstance(value, (list, tuple)) and len(value) == 3:
        return RGBColor(*(int(v) for v in value))
    raise TypeError(f"无法识别的颜色值：{value!r}（应为 'RRGGBB' 或 [r,g,b]）")


def assemble(spec: dict, spec_dir: Path, out: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    w_px = float(spec.get("width_px", 1920))
    emu_per_px = SLIDE_W_EMU / w_px
    pt_per_px = SLIDE_W_PT / w_px
    dy_map = spec.get("dy", {})
    paper = _rgb(spec.get("paper_color", "F7F0E4"))

    def set_font(run, latin, ea, size_px, color, spc, italic):
        f = run.font
        f.name = latin                      # 只写 a:latin
        f.size = Pt(size_px * pt_per_px)
        f.italic = italic
        f.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):        # 中日韩字形由 a:ea 决定，必须单独设
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", ea)
        if spc:
            rPr.set("spc", str(int(spc)))

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)

    for s_i, s_spec in enumerate(spec["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = paper

        # 图层顺序不可变：纸纹底 → 晕染层 → 全部原生文本框
        for role in ("paper", "wash"):
            rel = s_spec.get(role)
            if not rel:
                continue
            pic = slide.shapes.add_picture(str(spec_dir / rel), 0, 0,
                                           width=Emu(SLIDE_W_EMU),
                                           height=Emu(SLIDE_H_EMU))
            pic.name = {"paper": "paper-texture", "wash": "watercolor-layer"}[role]

        for blk in s_spec.get("blocks", []):
            left, top = blk["left"], blk["top"] + dy_map.get(blk["key"], 0.0)
            box = slide.shapes.add_textbox(
                Emu(_emu(left, emu_per_px)), Emu(_emu(top, emu_per_px)),
                Emu(_emu(blk["width"], emu_per_px)),
                Emu(_emu(blk["height"], emu_per_px)))
            tf = box.text_frame
            tf.word_wrap = False                    # wrap="none"，防意外折行
            tf.auto_size = MSO_AUTO_SIZE.NONE        # 关自动缩放，否则字号会被偷改
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = 0
            tf.margin_top = tf.margin_bottom = 0
            box.name = blk["key"]

            # 行距用精确值（spcPts）。倍数行距会乘上 CJK 字体 ~1.45em 的天然行高，位置全乱。
            pitch_px = blk.get("line_px") or blk["line_height"] * blk["size_px"]

            for p_i, runs in enumerate(blk["paras"]):
                p = tf.paragraphs[0] if p_i == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.LEFT
                p.line_spacing = Pt(pitch_px * pt_per_px)
                p.space_before = Pt(0)
                p.space_after = Pt(0)
                for item in runs:
                    ov = item[2] if len(item) > 2 else {}
                    run = p.add_run()
                    run.text = item[0]
                    set_font(run,
                             latin=ov.get("latin", blk["latin"]),
                             ea=ov.get("ea", blk["ea"]),
                             size_px=ov.get("size_px", blk["size_px"]),
                             color=_rgb(item[1]),
                             spc=ov.get("spc", blk.get("spc", 0)),
                             italic=ov.get("italic", blk.get("italic", False)))
        print(f"slide {s_i}: {len(s_spec.get('blocks', []))} 个文本框")

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"saved: {out}")


def cmd_assemble(argv: list[str]) -> int:
    if len(argv) != 2:
        print("用法：assemble <spec.json> <out.pptx>", file=sys.stderr)
        return 2
    spec_path = Path(argv[0]).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    assemble(spec, spec_path.parent, Path(argv[1]).resolve())
    return 0


# ===========================================================================
# measure · 逐元素测偏移
# ===========================================================================

def cmd_measure(argv: list[str]) -> int:
    if len(argv) != 3:
        print("用法：measure <ref_dir> <pptx_png_dir> <regions.json>", file=sys.stderr)
        return 2
    import numpy as np
    from PIL import Image

    ref_dir, ppt_dir = Path(argv[0]), Path(argv[1])
    regions = json.loads(Path(argv[2]).read_text(encoding="utf-8"))
    paper = np.array(regions.get("paper_rgb", [247, 240, 228]), dtype=np.int16)
    thresh = regions.get("threshold", 45)
    min_run = regions.get("min_run", 3)

    def ink_bbox(path: Path, box):
        x0, y0, x1, y1 = box
        a = np.asarray(Image.open(path).convert("RGB").crop((x0, y0, x1, y1)),
                       dtype=np.int16)
        mask = np.abs(a - paper).max(axis=2) > thresh
        rows = np.where(mask.sum(axis=1) >= min_run)[0]
        cols = np.where(mask.sum(axis=0) >= min_run)[0]
        if rows.size == 0 or cols.size == 0:
            return None
        return (x0 + int(cols[0]), y0 + int(rows[0]),
                x0 + int(cols[-1]), y0 + int(rows[-1]))

    print(f"{'元素':<14}{'HTML 墨水盒':<26}{'PPTX 墨水盒':<26}{'dx':>5}{'dy':>5}")
    print("-" * 84)
    worst = 0
    for e in regions["elements"]:
        b_ref = ink_bbox(ref_dir / e["ref"], e["box"])
        b_ppt = ink_bbox(ppt_dir / e["ppt"], e["box"])
        if b_ref is None or b_ppt is None:
            print(f"{e['name']:<14}-- 区域内没找到墨水，量测区域需避开晕染层 --")
            continue
        dx, dy = b_ppt[0] - b_ref[0], b_ppt[1] - b_ref[1]
        worst = max(worst, abs(dx), abs(dy))
        flag = "" if abs(dx) <= 2 and abs(dy) <= 2 else "   <-- 回填 dy: %+d" % (-dy)
        print(f"{e['name']:<14}{str(b_ref):<26}{str(b_ppt):<26}{dx:>5}{dy:>5}{flag}")
    print(f"\n最大偏移 {worst}px —— 目标 ≤1px；|dy|≤2 视为噪声，不要去「修正」")
    return 0


# ===========================================================================
# export · COM 回导
# ===========================================================================

def _new_engine():
    """起一个独占的 COM 实例。

    坑：版本无关的 `PowerPoint.Application` 在装了 WPS 的机器上会被 WPS 抢注，
    `Dispatch` 又偏爱复用已在跑的实例 —— 结果「点 MS 却拿到 WPS」。
    所以：① 用 `DispatchEx` 强制新实例；② ProgID 从版本号最高的开始试。
    详见 `10-drive-engine.md §四`。
    """
    import win32com.client
    last: Exception | None = None
    for progid in ("PowerPoint.Application.16", "PowerPoint.Application.15",
                   "PowerPoint.Application.14", "KWPP.Application",
                   "PowerPoint.Application"):
        try:
            return win32com.client.DispatchEx(progid)
        except Exception as exc:  # noqa: BLE001
            last = exc
    raise RuntimeError(f"无法启动 PowerPoint/WPS COM 引擎：{last}")


def cmd_export(argv: list[str]) -> int:
    if len(argv) < 2:
        print("用法：export <pptx> <out_dir> [width]", file=sys.stderr)
        return 2

    # ⚠️ 必须绝对路径：PowerPoint 会按自己的工作目录解析相对路径，报「找不到 slide01.png」
    pptx = str(Path(argv[0]).resolve())
    out_dir = Path(argv[1]).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    width = int(argv[2]) if len(argv) > 2 else 1920
    height = int(width * 9 / 16)

    app = _new_engine()
    pres = None
    try:
        pres = app.Presentations.Open(pptx, ReadOnly=True, Untitled=False,
                                      WithWindow=False)
        for i, slide in enumerate(pres.Slides, 1):
            dst = out_dir / f"slide{i:02d}.png"
            slide.Export(str(dst), "PNG", width, height)
            print("exported", dst.name, f"{width}x{height}")
        print("slides total:", pres.Slides.Count)
        return 0
    except Exception as exc:  # noqa: BLE001
        print("EXPORT FAILED:", exc, file=sys.stderr)
        return 1
    finally:
        if pres is not None:
            try:
                pres.Close()
            except Exception:
                pass
        try:
            app.Quit()
        except Exception:
            pass


COMMANDS = {
    "optimize": cmd_optimize,
    "assemble": cmd_assemble,
    "measure": cmd_measure,
    "export": cmd_export,
}


def main(argv: list[str]) -> int:
    if not argv or argv[0] in ("-h", "--help", "help"):
        print(__doc__)
        return 0
    if argv[0] not in COMMANDS:
        print(f"未知子命令 {argv[0]}；可用：{' · '.join(COMMANDS)}", file=sys.stderr)
        return 2
    return COMMANDS[argv[0]](argv[1:])


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
